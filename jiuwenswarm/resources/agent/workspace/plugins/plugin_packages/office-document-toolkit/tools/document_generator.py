import os
from pathlib import Path

from openjiuwen.core.foundation.tool import Tool, ToolCard

from pdf_font_utils import select_pdf_font


class DocumentGenerator(Tool):
    """文档生成工具：根据结构化内容生成 PDF、Word、Excel、PPT 文件。"""

    def __init__(self) -> None:
        super().__init__(
            ToolCard(
                id="document_generator",
                name="document_generator",
                description=(
                    "文档生成工具：根据结构化内容生成PDF、Word、Excel、PPT文件。"
                    "当用户需要创建文档、生成报告、输出文件时调用。"
                ),
                input_params={
                    "type": "object",
                    "properties": {
                        "format": {
                            "type": "string",
                            "enum": ["pdf", "word", "excel", "ppt"],
                            "description": "输出文件格式",
                        },
                        "filename": {
                            "type": "string",
                            "description": "输出文件名（不含扩展名）",
                        },
                        "content": {
                            "type": "object",
                            "description": (
                                "结构化内容，可包含 title, subtitle, paragraphs[], "
                                "tables[], sheets[], slides[]"
                            ),
                        },
                        "output_subdir": {
                            "type": "string",
                            "description": "输出子目录名，默认为 generated",
                        },
                    },
                    "required": ["format", "filename", "content"],
                },
            )
        )

    async def invoke(self, inputs, **kwargs):
        fmt = inputs.get("format", "")
        filename = inputs.get("filename", "")
        content = inputs.get("content", {})
        output_subdir = inputs.get("output_subdir", "generated")

        if not fmt or not filename or not content:
            return {
                "success": False,
                "error": "缺少必要参数: format, filename, content",
            }

        from openjiuwen.core.sys_operation.cwd import get_cwd

        base_dir = Path(get_cwd()) / output_subdir
        base_dir.mkdir(parents=True, exist_ok=True)

        ext_map = {
            "pdf": ".pdf",
            "word": ".docx",
            "excel": ".xlsx",
            "ppt": ".pptx",
        }
        ext = ext_map.get(fmt, f".{fmt}")
        output_path = base_dir / f"{filename}{ext}"

        try:
            if fmt == "pdf":
                self._generate_pdf(str(output_path), content)
            elif fmt == "word":
                self._generate_word(str(output_path), content)
            elif fmt == "excel":
                self._generate_excel(str(output_path), content)
            elif fmt == "ppt":
                self._generate_ppt(str(output_path), content)
            else:
                return {"success": False, "error": f"不支持的格式: {fmt}"}

            if not output_path.exists() or output_path.stat().st_size == 0:
                return {
                    "success": False,
                    "error": "文件生成失败，输出文件为空或不存在",
                }

            return {
                "success": True,
                "format": fmt,
                "path": str(output_path),
                "absolute_path": str(output_path.resolve()),
                "filename": output_path.name,
                "size_bytes": output_path.stat().st_size,
                "exists": True,
            }
        except ImportError as e:
            return {
                "success": False,
                "error": f"依赖库缺失: {e}. 请安装对应依赖后重试。",
            }
        except Exception as e:
            return {"success": False, "error": f"生成失败: {e}"}

    async def stream(self, inputs, **kwargs):
        yield await self.invoke(inputs, **kwargs)

    @staticmethod
    def _generate_pdf(file_path: str, content: dict) -> None:
        from fpdf import FPDF

        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)

        font_name = select_pdf_font(pdf)

        title = content.get("title", "")
        if title:
            pdf.set_font(font_name, "B", 16)
            pdf.multi_cell(0, 10, title)
            pdf.ln(5)

        pdf.set_font(font_name, "", 11)
        for para in content.get("paragraphs", []):
            text = para if isinstance(para, str) else para.get("text", "")
            if text:
                pdf.multi_cell(0, 7, text)
                pdf.ln(3)

        for table in content.get("tables", []):
            pdf.ln(5)
            data = table if isinstance(table, list) else table.get("data", [])
            if data:
                col_count = max(len(row) for row in data) if data else 1
                col_width = 180 / col_count
                for row in data:
                    for cell in row:
                        pdf.cell(col_width, 7, str(cell)[:50], border=1)
                    pdf.ln()

        pdf.output(file_path)

    @staticmethod
    def _generate_word(file_path: str, content: dict) -> None:
        from docx import Document

        doc = Document()

        title = content.get("title", "")
        if title:
            doc.add_heading(title, level=0)

        for para in content.get("paragraphs", []):
            text = para if isinstance(para, str) else para.get("text", "")
            style = (
                para.get("style", "Normal")
                if isinstance(para, dict)
                else "Normal"
            )
            if not text:
                continue
            if style.lower().startswith("heading"):
                num_part = "".join(c for c in style if c.isdigit())
                level = int(num_part) if num_part else 1
                doc.add_heading(text, level=level)
            else:
                doc.add_paragraph(text)

        for table in content.get("tables", []):
            data = table if isinstance(table, list) else table.get("data", [])
            if data:
                rows = len(data)
                cols = max(len(row) for row in data) if data else 1
                t = doc.add_table(rows=rows, cols=cols)
                for i, row in enumerate(data):
                    for j, cell in enumerate(row):
                        if j < cols:
                            t.rows[i].cells[j].text = str(cell)

        doc.save(file_path)

    @staticmethod
    def _generate_excel(file_path: str, content: dict) -> None:
        from openpyxl import Workbook

        wb = Workbook()
        sheets = content.get("sheets", [])
        if not sheets:
            ws = wb.active
            ws.title = content.get("sheet_name", "Sheet1")
            table_data = content.get("tables", [])
            if table_data:
                first_table = table_data[0]
                data = (
                    first_table
                    if isinstance(first_table, list)
                    else first_table.get("data", [])
                )
                for row in data:
                    ws.append(row)
            else:
                for row in content.get("rows", []):
                    ws.append(row)
        else:
            wb.remove(wb.active)
            for sheet_data in sheets:
                ws = wb.create_sheet(
                    title=sheet_data.get("sheet_name", "Sheet")
                )
                for row in sheet_data.get("rows", []):
                    ws.append(row)

        wb.save(file_path)

    @staticmethod
    def _generate_ppt(file_path: str, content: dict) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()

        title = content.get("title", "")
        if title:
            slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(slide_layout)
            if slide.shapes.title:
                slide.shapes.title.text = title
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = content.get("subtitle", "")

        for slide_data in content.get("slides", []):
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            shapes = slide.shapes

            slide_title = slide_data.get("title", "")
            if slide_title and shapes.title:
                shapes.title.text = slide_title

            body = slide_data.get("body", "")
            if body and len(shapes.placeholders) > 1:
                shapes.placeholders[1].text_frame.text = body

            for table_data in slide_data.get("tables", []):
                data = (
                    table_data
                    if isinstance(table_data, list)
                    else table_data.get("data", [])
                )
                if data:
                    rows = len(data)
                    cols = max(len(row) for row in data) if data else 1
                    table = shapes.add_table(
                        rows, cols, Inches(1), Inches(2), Inches(8), Inches(3)
                    ).table
                    for i, row in enumerate(data):
                        for j, cell in enumerate(row):
                            if j < cols:
                                table.cell(i, j).text = str(cell)

        prs.save(file_path)
