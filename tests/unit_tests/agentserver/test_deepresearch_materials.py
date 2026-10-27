"""Request material preparation and immutable snapshot contracts."""
import asyncio
import json
from pathlib import Path
import stat

import pytest
from jiuwenswarm.agents.harness.common.tools.deepresearch import materials as m


@pytest.mark.asyncio
@pytest.mark.parametrize('operation', ['fsync', 'replace'])
async def test_snapshot_publish_failure_removes_temporary_file(tmp_path, monkeypatch, operation):
    source = tmp_path / "source.txt"
    source.write_text("material content")
    context = {"original_query": "参考附件研究", "files": [{"path": str(source)}],
               "scope": "session", "output_dir": str(tmp_path)}
    failure = OSError("publish failed")

    def fail_replace(*args):
        raise failure

    monkeypatch.setattr(m.os, operation, fail_replace)
    with pytest.raises(OSError) as caught:
        await m.prepare_materials(context, "35c0b386-7e75-4923-b373-720ed2ed848b")
    assert caught.value is failure
    assert all(path.is_dir() for path in (tmp_path / ".deepresearch-materials").rglob("*"))


@pytest.mark.asyncio
async def test_snapshot_publish_closes_private_file_before_atomic_replace(tmp_path, monkeypatch):
    source = tmp_path / 'source.txt'
    source.write_text('material content')
    context = {'original_query': '参考附件研究', 'files': [{'path': str(source)}],
               'scope': 'session', 'output_dir': str(tmp_path)}
    original_replace = m.os.replace
    original_fsync = m.os.fsync
    descriptors = []

    def record_fsync(fd):
        original_fsync(fd)
        descriptors.append(fd)

    def checked_replace(temporary, target):
        assert len(descriptors) == 1
        with pytest.raises(OSError):
            m.os.fstat(descriptors[0])
        assert Path(temporary).parent == target.parent
        assert stat.S_IMODE(Path(temporary).stat().st_mode) == 0o600
        assert not target.exists()
        original_replace(temporary, target)

    monkeypatch.setattr(m.os, 'fsync', record_fsync)
    monkeypatch.setattr(m.os, 'replace', checked_replace)
    ref = await m.prepare_materials(context, '35c0b386-7e75-4923-b373-720ed2ed848b')
    target = Path(ref['path'])
    assert list(target.parent.iterdir()) == [target]
    assert stat.S_IMODE(target.stat().st_mode) == 0o600
    assert json.loads(target.read_text())['materials'][0]['content'] == 'material content'


@pytest.mark.asyncio
async def test_snapshot_large_text_partial_failure_and_dedup(tmp_path):
    source = tmp_path / "source.txt"
    source.write_text("独特事实青鹭项目效率为73.42%。" * 10000)
    context = {"original_query": "参考附件研究", "files": [
        {"path": str(source), "name": "source.txt"},
        {"path": str(source), "name": "same.txt"},
        {"path": str(tmp_path / "bad.txt"), "name": "bad.txt"}],
        "scope": "session-one", "output_dir": str(tmp_path)}
    ref = await m.prepare_materials(context, "35c0b386-7e75-4923-b373-720ed2ed848b")
    payload = json.loads(Path(ref["path"]).read_text())
    assert len(payload["materials"]) == 1
    assert len(payload["materials"][0]["content"]) > 65536
    assert [r["status"] for r in payload["results"]] == ["ready", "duplicate", "failed"]
    assert "content" not in ref


@pytest.mark.asyncio
async def test_required_failure_stops_and_network_prohibition(tmp_path):
    context = {"original_query": "必须依据附件研究", "files": [
        {"path": str(tmp_path / "missing.txt"), "name": "missing.txt"}],
        "scope": "s", "output_dir": str(tmp_path)}
    with pytest.raises(m.MaterialsError, match="核心材料"):
        await m.prepare_materials(context, "35c0b386-7e75-4923-b373-720ed2ed848b")
    context["original_query"] = "仅依据材料，禁止联网"
    with pytest.raises(m.MaterialsError, match="联网"):
        await m.prepare_materials(context, "35c0b386-7e75-4923-b373-720ed2ed848b")


def test_candidates_only_from_original_message():
    context = {"original_query": "参考 https://example.org/a 和 https://example.org/a。", "files": []}
    assert [c["source"] for c in m.candidates(context)] == ["https://example.org/a"]


@pytest.mark.parametrize("url", ["http://127.0.0.1", "https://169.254.169.254/a", "http://[::1]", "file:///etc/passwd", "http://user:pass@example.com"])
@pytest.mark.asyncio
async def test_unsafe_urls_rejected(url):
    with pytest.raises(m.MaterialsError):
        await m.fetch_url(url)


@pytest.mark.parametrize("query,required", [
    ("参考附件研究", False), ("必须依据附件研究", True), ("不必依据附件研究", False),
    ('讨论“必须依据附件”这句话', False), ('不是要求必须依据附件', False)])
def test_core_material_rules(query, required):
    rows = m.candidates({"original_query": query, "files": [{"path": "/file.txt", "name": "file.txt"}]})
    assert rows[0]["required"] is required


def test_sdk_policy_contract_and_no_material_compatibility():
    from openjiuwen_deepsearch.algorithm.query_understanding.material_processing import resolve_material_usage_mode
    assert resolve_material_usage_mode(m.MATERIALS_POLICY, True) == "required"
    assert m.context_from_request({"query": "研究主题"}, scope="s", output_dir="/tmp") is None
    context = m.context_from_request({"query": "generated https://not-source.test", "original_user_query": "参考 https://example.org"}, scope="s", output_dir="/tmp")
    assert [c["source"] for c in m.candidates(context)] == ["https://example.org"]


@pytest.mark.asyncio
async def test_limits_symlink_and_all_failed(tmp_path, monkeypatch):
    source = tmp_path / "source.txt"
    source.write_text('x' * 101)
    monkeypatch.setattr(m, 'MAX_TEXT', 100)
    context = {"original_query": "参考附件", "files": [{"path": str(source)}], "scope": "s", "output_dir": str(tmp_path)}
    ref = await m.prepare_materials(context, "35c0b386-7e75-4923-b373-720ed2ed848b")
    assert ref['results'][0]['status'] == 'limit_exceeded'
    assert '未获得可用材料' in m.material_summary(ref)
    link = tmp_path / 'link.txt'
    link.symlink_to(source)
    with pytest.raises(m.MaterialsError):
        await m.extract_attachment(str(link))
    monkeypatch.setattr(m, 'MAX_INPUT', 10)
    with pytest.raises(m.MaterialsError, match="限额"):
        await m.extract_attachment(str(source))


@pytest.mark.asyncio
async def test_docx_table_and_text_pdf_real_parsers(tmp_path):
    from docx import Document
    from PyPDF2 import PdfWriter
    from PyPDF2.generic import NameObject, DictionaryObject, DecodedStreamObject
    doc = Document()
    doc.add_paragraph('UNIQUE-FACT-7342')
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [('Metric', 'Value'), ('Efficiency', '73.42%')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    path = tmp_path / 'fixture.docx'
    doc.save(path)
    text = await m.extract_attachment(str(path))
    assert 'UNIQUE-FACT-7342' in text and 'Metric' in text and '73.42%' in text and ' | ' in text
    writer = PdfWriter()
    writer.add_blank_page(width=400, height=400)
    page = writer.pages[0]
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 300 Td (UNIQUE-PDF-7342) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    path = tmp_path / 'fixture.pdf'
    with path.open('wb') as output:
        writer.write(output)
    assert 'UNIQUE-PDF-7342' in await m.extract_attachment(str(path))


@pytest.mark.asyncio
async def test_dns_private_and_rebinding_rejected(monkeypatch):
    async def bad_dns(*args, **kwargs):
        return [(2, 1, 6, '', ('127.0.0.1', 80))]
    loop = asyncio.get_running_loop()
    monkeypatch.setattr(loop, 'getaddrinfo', bad_dns)
    with pytest.raises(Exception):
        await m.fetch_url('http://rebind.example.org')


@pytest.mark.asyncio
@pytest.mark.parametrize('proxy', [None, 'http://trusted-proxy.example:8080'])
@pytest.mark.parametrize('address', [
    '93.184.216.34', '198.18.0.0', '198.18.255.255', '198.19.0.0', '198.19.255.255',
])
async def test_domain_resolution_requires_public_addresses(monkeypatch, proxy, address):
    import aiohttp
    import ssl
    from jiuwenswarm.agents.harness.common.tools import web_fetch_tools

    monkeypatch.setattr(web_fetch_tools, '_get_free_search_proxy_url', lambda: proxy or '')

    async def fake_dns(*args, **kwargs):
        return [(2, 1, 6, '', (address, 443))]

    monkeypatch.setattr(asyncio.get_running_loop(), 'getaddrinfo', fake_dns)

    class Content:
        async def iter_chunked(self, size):
            yield b'<html><title>Queues</title><p>Queue maxsize applies backpressure.</p></html>'

    class Response:
        status = 200
        headers = {'Content-Type': 'text/html'}
        content_length = None
        charset = 'utf-8'
        content = Content()

        def __init__(self, resolver):
            self.resolver = resolver

        async def __aenter__(self):
            records = await self.resolver.resolve('docs.python.org', 443)
            assert records[0]['host'] == address
            return self

        async def __aexit__(self, *args):
            return None

    class Client:
        def __init__(self, **kwargs):
            self.connector = kwargs['connector']
            assert self.connector._ssl.verify_mode == ssl.CERT_REQUIRED

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            await self.connector.close()

        def get(self, url, **kwargs):
            assert kwargs['proxy'] == proxy
            return Response(self.connector._resolver)

    monkeypatch.setattr(aiohttp, 'ClientSession', Client)
    if address != '93.184.216.34':
        with pytest.raises(m.MaterialsPolicyError, match='网页地址不是公开网络地址'):
            await m.fetch_url('https://docs.python.org/3/library/asyncio-queue.html')
        return
    text, title, metadata = await m.fetch_url('https://docs.python.org/3/library/asyncio-queue.html')
    assert 'backpressure' in text and title == 'Queues'
    assert metadata['original_url'] == metadata['final_url']


@pytest.mark.parametrize('url', ['http://198.18.2.84/', 'https://198.19.255.255/'])
def test_literal_fake_ip_remains_rejected(url):
    with pytest.raises(m.MaterialsError):
        m._validate_url(url)


@pytest.mark.asyncio
async def test_previous_proxy_host_cannot_bypass_direct_rebinding_check(monkeypatch):
    import aiohttp
    from jiuwenswarm.agents.harness.common.tools import web_fetch_tools

    proxy_host = 'trusted-proxy.example'
    monkeypatch.setattr(web_fetch_tools, '_get_free_search_proxy_url', lambda: f'http://{proxy_host}:8080')
    monkeypatch.setenv('NO_PROXY', proxy_host)
    target_resolutions = 0

    async def changing_dns(host, port, **kwargs):
        nonlocal target_resolutions
        address = '8.8.8.8'
        if host == proxy_host and port == 80:
            target_resolutions += 1
            if target_resolutions > 1:
                address = '127.0.0.1'
        return [(2, 1, 6, '', (address, port))]

    monkeypatch.setattr(asyncio.get_running_loop(), 'getaddrinfo', changing_dns)

    class Response:
        status = 302
        headers = {'Location': f'http://{proxy_host}/source'}

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return None

    class Client:
        def __init__(self, **kwargs):
            self.connector = kwargs['connector']

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            await self.connector.close()

        def get(self, url, **kwargs):
            if url.startswith('http://'):
                assert kwargs['proxy'] is None

                class DirectResponse(Response):
                    async def __aenter__(inner_self):
                        await self.connector._resolver.resolve(proxy_host, 80)
                        raise AssertionError('Direct connection accepted a private rebound address')

                return DirectResponse()
            return Response()

    monkeypatch.setattr(aiohttp, 'ClientSession', Client)
    with pytest.raises(m.MaterialsError, match='公开网络地址'):
        await m.fetch_url('https://public.example.org/start')

@pytest.mark.parametrize('url', ['http://127.1', 'http://2130706433', 'http://0177.0.0.1'])
def test_numeric_host_cannot_bypass_connector_resolver(url):
    with pytest.raises(m.MaterialsError):
        m._validate_url(url)


@pytest.mark.parametrize('query', ['请研究附件，不要联网，只根据附件完成报告', '不要联网', '不用联网研究'])
def test_offline_requirements_are_not_removed_by_negation(query):
    assert m.context_from_request({'query': query}, scope='s', output_dir='/tmp') is not None


@pytest.mark.parametrize('query', ['必须依据附件研究，但不要生成简报', '必须使用核心材料 a.txt，其他材料不用引用'])
def test_unrelated_negation_does_not_remove_core_requirement(query):
    assert m.candidates({'original_query': query, 'files': [{'path': '/a.txt', 'name': 'a.txt'}]})[0]['required']


@pytest.mark.asyncio
async def test_large_docx_is_read_to_eof(tmp_path):
    from docx import Document
    doc = Document()
    content = 'FACT-7342 ' * 20000
    doc.add_paragraph(content)
    path = tmp_path / 'large.docx'
    doc.save(path)
    result = await m.extract_attachment(str(path))
    assert result.strip() == content.strip()


@pytest.mark.asyncio
async def test_cancel_parser_with_paused_output_reaps_child(tmp_path, monkeypatch):
    import sys
    source = tmp_path / 'cancel.pdf'
    source.write_bytes(b'fixture')
    original_spawn = asyncio.create_subprocess_exec
    entered = asyncio.Event()
    children = []
    async def spawn(*args, **kwargs):
        process = await original_spawn(sys.executable, '-c',
            'import sys; sys.stdin.buffer.read(); sys.stdout.buffer.write(b"x" * 1000000); sys.stdout.flush()', **kwargs)
        children.append(process)
        original_read = process.stdout.read
        reads = 0
        async def paused_first_read(size=-1):
            nonlocal reads
            reads += 1
            if reads == 1:
                await asyncio.sleep(0.05)
                entered.set()
                await asyncio.Event().wait()
            return await original_read(size)
        process.stdout.read = paused_first_read
        return process
    monkeypatch.setattr(m.asyncio, 'create_subprocess_exec', spawn)
    task = asyncio.create_task(m.extract_attachment(str(source)))
    await asyncio.wait_for(entered.wait(), 3)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await asyncio.wait_for(task, 3)
    assert children[0].returncode is not None

@pytest.mark.asyncio
async def test_url_redirect_stream_limits_and_html_table(monkeypatch):
    import aiohttp
    async def public_dns(*args, **kwargs):
        return [(2, 1, 6, '', ('8.8.8.8', 443))]
    monkeypatch.setattr(asyncio.get_running_loop(), 'getaddrinfo', public_dns)
    hits, responses = [], []
    class Content:
        async def iter_chunked(self, size):
            yield b'<html><title>Fixture</title><table><tr><th>Metric</th><th>Value</th></tr><tr><td>HERON</td><td>73.42</td></tr></table><a href="http://private.invalid">do not follow</a></html>'
    class Response:
        def __init__(self, status=200, headers=None):
            self.status = status
            self.headers = headers or {'Content-Type': 'text/html'}
            self.content_length = None
            self.charset = 'utf-8'
            self.content = Content()
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return None
    class Client:
        def __init__(self, **kwargs):
            assert kwargs['trust_env'] is False and kwargs['auto_decompress'] is False
            self.connector = kwargs['connector']
        async def __aenter__(self): return self
        async def __aexit__(self, *args): await self.connector.close()
        def get(self, url, **kwargs):
            assert kwargs['allow_redirects'] is False
            hits.append(url)
            return responses.pop(0)
    monkeypatch.setattr(aiohttp, 'ClientSession', Client)
    responses.append(Response())
    text, title, metadata = await m.fetch_url('https://public.example.org')
    assert 'Metric | Value' in text and 'HERON | 73.42' in text
    assert title == 'Fixture' and metadata['original_url'] == metadata['final_url']
    assert len(hits) == 1
    responses.append(Response(302, {'Location': 'http://127.1/metadata'}))
    with pytest.raises(m.MaterialsError):
        await m.fetch_url('https://public.example.org')
    assert len(hits) == 2  # rejected redirect never reaches HTTP client
    responses.append(Response())
    monkeypatch.setattr(m, 'MAX_INPUT', 10)
    with pytest.raises(m.MaterialsError, match='限额'):
        await m.fetch_url('https://public.example.org')


@pytest.mark.asyncio
async def test_candidate_count_and_total_content_limit(tmp_path, monkeypatch):
    paths = []
    for index in range(12):
        path = tmp_path / f'{index}.txt'
        path.write_text(str(index) * 20)
        paths.append({'path': str(path)})
    context = {'original_query': '参考附件', 'files': paths, 'scope': 's', 'output_dir': str(tmp_path)}
    monkeypatch.setattr(m, 'MAX_TOTAL', 50)
    reference = await m.prepare_materials(context, '35c0b386-7e75-4923-b373-720ed2ed848b')
    assert len(reference['results']) == 12
    assert reference['results'][10]['status'] == 'limit_exceeded'
    body = json.loads(Path(reference['path']).read_text())
    assert sum(len(item['content']) for item in body['materials']) <= 50


@pytest.mark.asyncio
async def test_utf16_bom_and_malformed_optional_url(tmp_path):
    source = tmp_path / 'windows.txt'
    source.write_bytes('中文资料7342'.encode('utf-16'))
    assert await m.extract_attachment(str(source)) == '中文资料7342'
    context = m.context_from_request({'original_user_query': '参考 https://[bad/path'}, scope='s', output_dir=str(tmp_path))
    reference = await m.prepare_materials(context, '35c0b386-7e75-4923-b373-720ed2ed848b')
    assert reference['results'][0]['status'] == 'failed'
